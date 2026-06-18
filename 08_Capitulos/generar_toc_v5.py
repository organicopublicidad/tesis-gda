from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Márgenes
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

AZUL = RGBColor(0x1F, 0x49, 0x7D)
NARANJA = RGBColor(0xC0, 0x50, 0x00)

def titulo_cap(doc, num, title, nota=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"CAPÍTULO {num}: {title}")
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = AZUL
    if nota:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(nota)
        r2.italic = True; r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(0x70,0x70,0x70)

def bloque(doc, items):
    for num, title, subs in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{num}. {title}")
        r.bold = True; r.font.size = Pt(10)
        for s in subs:
            p2 = doc.add_paragraph()
            p2.paragraph_format.left_indent = Cm(0.7)
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(s)
            r2.font.size = Pt(9)
            if "NUEVO" in s:
                r2.font.color.rgb = NARANJA

# ========== PORTADA ==========
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TABLA DE CONTENIDO — PLAN DE INVESTIGACIÓN")
r.bold = True; r.font.size = Pt(14); r.font.color.rgb = AZUL
p.paragraph_format.space_after = Pt(8)

meta = [
    ("Título:", "Gobernanza Digital Asociativa: un modelo sociotécnico para la captura de valor en origen en organizaciones caficultoras de territorios PDET de Colombia"),
    ("Programa:", "Doctorado TAPAS — ETSIAAB, Universidad Politécnica de Madrid"),
    ("Directora:", "Dra. Sonia Benito Hernández"),
    ("Candidato:", "Daniel Alejandro Pajoy Bastos"),
    ("Versión:", "0.5 — 17 de junio de 2026"),
]
for label, value in meta:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + " "); r.bold = True; r.font.size = Pt(11)
    r2 = p.add_run(value); r2.font.size = Pt(11)

doc.add_paragraph()
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Pregunta de investigación: "); r.bold = True; r.font.size = Pt(11)
r2 = p.add_run("¿Cómo un modelo sociotécnico de gobernanza digital asociativa habilita a las organizaciones caficultoras de pequeña escala en territorios PDET de Colombia para incrementar la proporción de valor capturado en origen (VCO) en la cadena global de valor del café?")
r2.italic = True; r2.font.size = Pt(11)

doc.add_paragraph("─" * 80)

# ========== OBJETIVOS ==========
p = doc.add_paragraph()
r = p.add_run("OBJETIVOS DE LA INVESTIGACIÓN")
r.bold = True; r.font.size = Pt(12); r.font.color.rgb = AZUL
p.paragraph_format.space_before = Pt(6)

p = doc.add_paragraph()
r = p.add_run("Objetivo general (OG): "); r.bold = True
p.add_run("Desarrollar y evaluar un modelo sociotécnico de gobernanza digital asociativa (GDA) que habilite a las organizaciones caficultoras de pequeña escala en territorios PDET de Colombia para incrementar la proporción de valor capturado en origen (VCO) en la cadena global de valor del café.")
p.paragraph_format.space_after = Pt(6)

oes = [
    ("OE1", "I — Diagnóstica", "Diagnosticar la distribución del valor a lo largo de la cadena del café colombiano y establecer la línea base de VCO₀ y del capital social de las organizaciones caficultoras en territorios PDET, construyendo el IPDET como moderador e identificando mediante fsQCA las configuraciones institucionales necesarias y suficientes para VCO alto.", "Cap. 6"),
    ("OE2", "II — Constructiva", "Diseñar de forma participativa, mediante Design Science Research y un panel Delphi, el modelo GDA que articula el subsistema social (gobernanza asociativa, Ostrom) y el subsistema técnico (trazabilidad y soberanía de datos) bajo el principio de optimización conjunta.", "Cap. 7"),
    ("OE3", "II → III", "Operacionalizar y parametrizar los constructos del modelo (Índice de Gobernanza, componentes técnicos y VCO) y construir las distribuciones a priori, a partir del consenso Delphi, que alimentan la simulación.", "Caps. 7-8"),
    ("OE4", "III — Evaluativa", "Evaluar el efecto del modelo GDA sobre el VCO y su robustez mediante simulación Monte Carlo, índices de Sobol y análisis de escenarios adversos, contrastando la hipótesis de que la gobernanza explica más varianza del VCO que la tecnología.", "Cap. 8"),
]
tabla = doc.add_table(rows=1, cols=4)
tabla.style = "Table Grid"
for i, txt in enumerate(["OE", "Fase", "Enunciado", "Capítulo"]):
    c = tabla.rows[0].cells[i]
    c.text = txt
    c.paragraphs[0].runs[0].bold = True
    c.paragraphs[0].runs[0].font.size = Pt(10)
for oe, fase, enunciado, cap in oes:
    row = tabla.add_row().cells
    for i, txt in enumerate([oe, fase, enunciado, cap]):
        row[i].text = txt
        row[i].paragraphs[0].runs[0].font.size = Pt(9)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
r = p.add_run("Hipótesis central: "); r.bold = True
p.add_run("La dimensión de gobernanza (subsistema social) explica una mayor proporción de la varianza del VCO que la dimensión tecnológica (subsistema técnico), evidenciando que la optimización conjunta —no la mera digitalización— es la condición habilitante. Se contrasta con la descomposición de varianza de Sobol.")
p.paragraph_format.space_after = Pt(10)
doc.add_paragraph("─" * 80)

# ========== PREFACIO ==========
for txt in ["ABSTRACT (Inglés)", "RESUMEN (Español)", "LISTA DE TABLAS", "LISTA DE FIGURAS", "LISTA DE ACRÓNIMOS Y ABREVIATURAS"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt)
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x40,0x40,0x40)

doc.add_paragraph()

# ========== CAPÍTULOS ==========

titulo_cap(doc, 1, "INTRODUCCIÓN A LA TESIS")
bloque(doc, [
    ("1.1", "Introducción", []),
    ("1.2", "Contexto de la investigación", ["1.2.1. El café en la economía colombiana","1.2.2. Distribución del valor en la cadena global y la brecha del VCO","1.2.3. El minifundio cafetero como universo de estudio"]),
    ("1.3", "El posconflicto colombiano y los PDET", ["1.3.1. El Acuerdo de Paz de 2016","1.3.2. Los Programas de Desarrollo con Enfoque Territorial","1.3.3. La caficultura como vehículo de reincorporación"]),
    ("1.4", "La gobernanza como factor habilitante", ["1.4.1. Gobernanza de comunes en organizaciones productoras","1.4.2. La transformación digital de las organizaciones rurales","1.4.3. La sostenibilidad como atributo de valor"]),
    ("1.5", "Planteamiento del problema de investigación", ["1.5.1. La brecha del valor capturado en origen (VCO)","1.5.2. Los cuatro déficits estructurales","1.5.3. Pregunta de investigación, hipótesis y objetivos"]),
    ("1.6", "Visión general de la metodología", [
        "1.6.1. Design Science Research (Hevner et al., 2004): tres fases",
        "1.6.2. Análisis Cualitativo Comparativo (fsQCA): censo FEMNCAFE (~20)  ◄ NUEVO",
        "1.6.3. Simulación Monte Carlo e índices de sensibilidad de Sobol",
        "1.6.4. Enfoque mixto secuencial: integración de métodos cualitativos y cuantitativos",
    ]),
    ("1.7", "Contribuciones esperadas", ["1.7.1. Contribución teórica","1.7.2. Contribución metodológica","1.7.3. Contribución praxeológica"]),
    ("1.8", "Estructura de la tesis", []),
])

titulo_cap(doc, 2, "EL CAFÉ EN COLOMBIA Y EN EL MUNDO")
bloque(doc, [
    ("2.1", "Introducción", []),
    ("2.2", "Historia del café en Colombia", []),
    ("2.3", "Tendencias mundiales en la producción", ["2.3.1. Principales países productores","2.3.2. Posición de Colombia en el mercado mundial","2.3.3. Evolución de la productividad"]),
    ("2.4", "Consumo, importaciones y exportaciones", ["2.4.1. Principales mercados de destino del café colombiano","2.4.2. Estados Unidos como mercado principal","2.4.3. La Unión Europea como mercado estratégico"]),
    ("2.5", "Precios del café", ["2.5.1. El Contrato C y la formación internacional de precios","2.5.2. El precio interno colombiano y el FEPCafé","2.5.3. Diferencial UGQ y prima de calidad"]),
    ("2.6", "La cadena de valor del café colombiano", ["2.6.1. Eslabones productivos y comerciales","2.6.2. Distribución del valor entre actores","2.6.3. El caficultor en el eslabón de menor captura"]),
    ("2.7", "Segmentación del café", ["2.7.1. Café convencional vs. cafés especiales","2.7.2. Métodos de procesamiento (lavado, honey, natural)","2.7.3. Café de origen y trazabilidad"]),
    ("2.8", "Políticas e institucionalidad cafetera en Colombia", ["2.8.1. La Federación Nacional de Cafeteros (FNC)","2.8.2. El FEPCafé y el costo medio de producción","2.8.3. Programas de apoyo y financiamiento"]),
    ("2.9", "Caficultura y posconflicto", ["2.9.1. Territorios PDET cafeteros","2.9.2. Iniciativas de café para la paz","2.9.3. La red FEMNCAFE"]),
])

titulo_cap(doc, 3, "GOBERNANZA Y CADENAS GLOBALES DE VALOR")
bloque(doc, [
    ("3.1", "Introducción", []),
    ("3.2", "El marco analítico de las Cadenas Globales de Valor", ["3.2.1. Origen y evolución del concepto","3.2.2. Las tres variables determinantes de Gereffi, Humphrey y Sturgeon","3.2.3. Los cinco tipos de gobernanza"]),
    ("3.3", "La cadena cautiva del café colombiano", ["3.3.1. Caracterización como oligopsonio","3.3.2. Mecanismos de extracción de renta territorial","3.3.3. Asimetrías estructurales de poder"]),
    ("3.4", "La evolución del poder en la cadena del café (Grabs y Ponte, 2019)", ["3.4.1. Tres fases históricas (pre-1989, 1989-2000, post-2000)","3.4.2. Los cuatro tipos de poder en la cadena","3.4.3. La paradoja de los sellos de sostenibilidad"]),
    ("3.5", "Upgrading en cadenas globales", ["3.5.1. Upgrading de proceso, producto, funcional e intersectorial","3.5.2. Trayectorias posibles para organizaciones de pequeños productores","3.5.3. Hacia cadenas modulares y relacionales"]),
    ("3.6", "Redes de producción global (GPN 2.0)", ["3.6.1. Actores no empresariales en la cadena","3.6.2. El papel del Estado, las ONG y la cooperación internacional","3.6.3. Aplicación al contexto PDET"]),
    ("3.7", "Diagnóstico del café colombiano como caso paradigmático", []),
])

titulo_cap(doc, 4, "GOBERNANZA DE COMUNES, POSCONFLICTO Y SOSTENIBILIDAD")
bloque(doc, [
    ("4.1", "Introducción", []),
    ("4.2", "La teoría de gobernanza de comunes (Ostrom, 1990)", ["4.2.1. El problema de los recursos de uso común","4.2.2. Los ocho principios de diseño institucional","4.2.3. Validación empírica multicontextual (Agrawal, 2001)"]),
    ("4.3", "Capital social y posconflicto en Colombia", ["4.3.1. Erosión del capital social por el conflicto armado (Rettberg, 2010)","4.3.2. La cultura colectiva de los firmantes de paz como activo","4.3.3. Ausencia estatal histórica como causa estructural"]),
    ("4.4", "Cooperativismo agroalimentario y sus fallas (Ortmann y King, 2007)", ["4.4.1. Estructuras de gobernanza en cooperativas cafeteras colombianas","4.4.2. Análisis institucional de la red FEMNCAFE"]),
    ("4.5", "Adaptación de Ostrom al posconflicto", ["4.5.1. Reformulación contextual de los ocho principios","4.5.2. Principios complementarios para alta fragilidad institucional","4.5.3. Operacionalización: el Índice de Gobernanza (IG)"]),
    ("4.6", "Investigación acción participativa (Fals Borda, 1986)", ["4.6.1. Fundamento epistemológico para la co-construcción","4.6.2. Implicaciones metodológicas"]),
    ("4.7", "Sostenibilidad y compromiso ambiental en PYMES", ["4.7.1. Apoyo público e inversión ambiental (Benito-Hernández et al., 2023)","4.7.2. PDET como instrumento de política habilitadora"]),
    ("4.8", "Valor implícito de los atributos de sostenibilidad en el café", ["4.8.1. El modelo de precios hedónicos (Merbah & Benito-Hernández, 2023)","4.8.2. La prima por trazabilidad verificada","4.8.3. Implicaciones para el VCO"]),
])

titulo_cap(doc, 5, "TECNOLOGÍA DIGITAL, SOBERANÍA DE DATOS Y DISEÑO SOCIOTÉCNICO")
bloque(doc, [
    ("5.1", "Introducción", []),
    ("5.2", "La teoría de sistemas sociotécnicos", ["5.2.1. Orígenes: la tradición de Tavistock (Trist y Bamforth, 1951)","5.2.2. Sistemas abiertos y optimización conjunta (Emery y Trist, 1965; Trist, 1981)","5.2.3. Vigencia contemporánea (Abbas y Michael, 2026; Horváth y Kemser, 2026)","5.2.4. El subsistema social y el subsistema técnico del GDA"]),
    ("5.3", "Blockchain y gobernanza de comunes", ["5.3.1. Las seis affordances de Rozas et al. (2021)","5.3.2. Codificación de la confianza","5.3.3. Riesgos de la adopción top-down"]),
    ("5.4", "Blockchain en cadenas agroalimentarias", ["5.4.1. Trazabilidad y verificación (Werbach, 2018)","5.4.2. Aplicaciones en café (Miatton y Amado, 2020; Elmatsani et al., 2026)","5.4.3. El EUDR y la georreferenciación obligatoria"]),
    ("5.5", "Soberanía de datos como recurso de uso común", ["5.5.1. Los datos de la finca como activo colectivo","5.5.2. Arquitecturas técnicas de soberanía","5.5.3. Implicaciones éticas"]),
    ("5.6", "Instrumentos financieros complementarios al VCO: blended finance y créditos de carbono", ["5.6.1. Blended finance para agricultura (Havemann et al., 2020)","5.6.2. Blockchain y regenerative finance (Schletz et al., 2023)","5.6.3. Instrumentos por horizonte temporal e implicaciones para la captura de valor en origen"]),
    ("5.7", "Síntesis del marco teórico integrado", ["5.7.1. Articulación de los pilares: CGV + comunes + digital bajo la lente sociotécnica","5.7.2. El modelo conceptual del GDA como sistema sociotécnico de optimización conjunta","5.7.3. Proposiciones teóricas e hipótesis derivadas"]),
])

titulo_cap(doc, 6, "DIAGNÓSTICO DEL VALOR CAPTURADO EN ORIGEN — ANÁLISIS DE LA CADENA DEL CAFÉ COLOMBIANO PDET", "(Fase I — Diagnóstica · Artículo 1)")
bloque(doc, [
    ("6.1", "Introducción", []),
    ("6.2", "Marco analítico de la distribución de valor", ["6.2.1. Definición y medición del Valor Capturado en Origen (VCO)","6.2.2. Variables del modelo de diagnóstico","6.2.3. Construcción del Índice de Consolidación PDET (IPDET) como moderador"]),
    ("6.3", "Metodología", [
        "6.3.1. Enfoque mixto secuencial y rol del DSR en la Fase I",
        "6.3.2. Selección de organizaciones y territorios (censo FEMNCAFE ~20; caso de validación: Cauca)",
        "6.3.3. Instrumentos cualitativos: entrevistas semiestructuradas, SNA y mapeo de la cadena",
        "6.3.4. Análisis Cualitativo Comparativo (fsQCA): configuraciones institucionales para VCO alto  ◄ NUEVO",
        "6.3.5. Integración de resultados cualitativos y cuantitativos",
    ]),
    ("6.4", "Resultados", [
        "6.4.1. Mapeo de la cadena del café colombiano hasta el mercado europeo",
        "6.4.2. Distribución cuantitativa del valor por eslabón",
        "6.4.3. VCO inicial (VCO₀) de las organizaciones estudiadas",
        "6.4.4. Diagnóstico del capital social posconflicto (resultados SNA)",
        "6.4.5. IPDET por municipio y categorización territorial",
        "6.4.6. Configuraciones necesarias y suficientes para VCO alto (resultados fsQCA)  ◄ NUEVO",
    ]),
    ("6.5", "Discusión", ["6.5.1. Comparación con la literatura previa","6.5.2. Patrones emergentes en la red FEMNCAFE","6.5.3. Implicaciones para el diseño del modelo GDA"]),
])

titulo_cap(doc, 7, "DISEÑO Y VALIDACIÓN DEL MODELO GDA — ENFOQUE PARTICIPATIVO Y DELPHI", "(Fase II — Constructiva · Artículo 2)")
bloque(doc, [
    ("7.1", "Introducción", []),
    ("7.2", "Diseño científico de artefactos (Design Science Research)", ["7.2.1. El paradigma de Hevner et al. (2004)","7.2.2. Las tres condiciones de validez","7.2.3. Aplicación al GDA"]),
    ("7.3", "Metodología", ["7.3.1. Talleres de diseño participativo con las organizaciones","7.3.2. Panel Delphi: composición experta y protocolo de las tres rondas","7.3.3. Validación comunitaria"]),
    ("7.4", "Resultados", ["7.4.1. Co-diseño de los mecanismos de gobernanza (subsistema social)","7.4.2. Arquitectura técnica del componente blockchain (subsistema técnico)","7.4.3. Consenso experto Delphi por variable","7.4.4. Distribuciones a priori construidas para la simulación","7.4.5. Prototipo del modelo GDA validado"]),
    ("7.5", "Discusión", ["7.5.1. El modelo GDA como artefacto sociotécnico de optimización conjunta","7.5.2. Soberanía de datos como innovación institucional","7.5.3. Diferenciación frente a modelos previos"]),
])

titulo_cap(doc, 8, "EVALUACIÓN DEL MODELO — MONTE CARLO, SOBOL Y PRUEBAS DE ESTRÉS", "(Fase III — Evaluativa · Artículo 3)")
bloque(doc, [
    ("8.1", "Introducción", []),
    ("8.2", "Marco evaluativo", ["8.2.1. Hipótesis a contrastar (gobernanza explica más varianza del VCO que la tecnología)","8.2.2. Variables, parámetros e indicadores","8.2.3. Indicadores de robustez del VCO"]),
    ("8.3", "Metodología", ["8.3.1. Configuración del modelo de simulación Monte Carlo","8.3.2. Implementación de los índices de sensibilidad de Sobol","8.3.3. Diseño de escenarios adversos (pruebas de estrés)","8.3.4. Piloto acotado e indicadores tempranos de adopción"]),
    ("8.4", "Resultados", [
        "8.4.1. Distribución del VCO bajo el modelo GDA (10.000 iteraciones)",
        "8.4.2. Descomposición de varianza: gobernanza vs. tecnología (índices de Sobol)",
        "8.4.3. Respuesta del VCO bajo escenarios adversos (pruebas de estrés)",
        "8.4.4. Robustez de los resultados",
        "8.4.5. Resultados del piloto acotado",
        "8.4.6. Transferibilidad en la red FEMNCAFE (validez externa vía IPDET)",
    ]),
    ("8.5", "Discusión", ["8.5.1. Validación o refutación de la hipótesis central","8.5.2. Alcance y límites de la robustez de los hallazgos","8.5.3. Limitaciones del modelo"]),
])

titulo_cap(doc, 9, "CONCLUSIÓN GENERAL")
bloque(doc, [
    ("9.1", "Síntesis de hallazgos por capítulo", []),
    ("9.2", "Respuesta a la pregunta de investigación", []),
    ("9.3", "Validación de las hipótesis y cumplimiento de los objetivos", []),
    ("9.4", "Contribuciones originales al conocimiento", ["9.4.1. Aportes teóricos","9.4.2. Aportes metodológicos","9.4.3. Aportes praxeológicos"]),
    ("9.5", "Implicaciones para la política pública", ["9.5.1. Para el marco PDET","9.5.2. Para la política cafetera colombiana","9.5.3. Para la regulación europea de sostenibilidad (EUDR)"]),
    ("9.6", "Implicaciones prácticas para las organizaciones", []),
    ("9.7", "Limitaciones del estudio", []),
    ("9.8", "Líneas futuras de investigación", []),
])

# ========== COLA ==========
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("ANEXOS")
r.bold = True; r.font.size = Pt(12); r.font.color.rgb = AZUL
p.paragraph_format.space_before = Pt(10)

anexos = [
    ("Anexo 1", "Instrumentos de campo", ["A1.1. Guion de entrevistas semiestructuradas","A1.2. Protocolo del Análisis de Redes Sociales (SNA)","A1.3. Cuestionario del panel Delphi","A1.4. Formularios de consentimiento informado","A1.5. Encuesta de transferibilidad FEMNCAFE"]),
    ("Anexo 2", "Operacionalización de constructos", ["A2.1. Índice de Gobernanza (IG): indicadores y escalas de calibración","A2.2. IPDET: dimensiones e indicadores","A2.3. Cálculo del VCO: ecuaciones y fuentes de datos","A2.4. Tablas de calibración fsQCA  ◄ NUEVO"]),
    ("Anexo 3", "Implementación computacional", ["A3.1. Código de la simulación Monte Carlo","A3.2. Implementación de los índices de Sobol","A3.3. Diseño de escenarios adversos (pruebas de estrés)","A3.4. Validación del modelo"]),
    ("Anexo 4", "Aprobación ética", ["A4.1. Dictamen del comité de ética institucional","A4.2. Convenios con organizaciones participantes"]),
    ("Anexo 5", "Publicaciones derivadas", ["A5.1. Artículo 1: Diagnóstico de la distribución de valor (Fase I)","A5.2. Artículo 2: El modelo de gobernanza GDA (Fase II)","A5.3. Artículo 3: Evaluación bajo incertidumbre (Fase III)"]),
]
for label, title, subs in anexos:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{label}. {title}"); r.bold = True; r.font.size = Pt(10)
    for s in subs:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.7)
        p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(s); r2.font.size = Pt(9)
        if "NUEVO" in s:
            r2.font.color.rgb = NARANJA

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("REFERENCIAS"); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = AZUL
doc.add_paragraph("(En orden alfabético, formato APA 7, con DOI cuando esté disponible)").runs[0].italic = True

doc.add_paragraph("─" * 80)

# RD 99/2011
p = doc.add_paragraph()
r = p.add_run("MEDIOS Y RECURSOS (RD 99/2011)")
r.bold = True; r.font.size = Pt(11); r.font.color.rgb = AZUL
p.paragraph_format.space_before = Pt(10)
medios = [
    "Acceso al terreno: vínculo operativo con COFFIS y la red FEMNCAFE (~20 organizaciones — censo completo). Sede del candidato en Popayán, Cauca (caso de validación).",
    "Datos: entrevistas semiestructuradas, datos de cadena por eslabón, datos de finca; panel de expertos para el Delphi; condiciones institucionales para el fsQCA.",
    "Recursos computacionales: Python (Monte Carlo, Sobol, 10.000 iteraciones); software fsQCA 3.0 / R-QCA; herramientas SNA.",
    "Dirección y marco institucional: Dra. Sonia Benito (ETSIAAB-UPM, grupo BIDA); programa TAPAS; modalidad virtual desde Colombia con trabajo de campo in situ.",
    "Ética: dictamen del comité de ética institucional y convenios de consentimiento informado (Anexo 4).",
]
for m in medios:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("• " + m); r.font.size = Pt(10)

p = doc.add_paragraph()
r = p.add_run("PLANIFICACIÓN TEMPORAL (3–4 años)")
r.bold = True; r.font.size = Pt(11); r.font.color.rgb = AZUL
p.paragraph_format.space_before = Pt(8)
cron = doc.add_table(rows=1, cols=3)
cron.style = "Table Grid"
for i, txt in enumerate(["Periodo", "Fase / Hito", "Producto"]):
    c = cron.rows[0].cells[i]
    c.text = txt
    c.paragraphs[0].runs[0].bold = True
    c.paragraphs[0].runs[0].font.size = Pt(10)
for periodo, fase, prod in [
    ("Año 1", "Marco teórico (Caps. 1-5) + Fase I diagnóstica: trabajo de campo, SNA, fsQCA, VCO₀, IPDET", "CAPD · Artículo 1"),
    ("Año 2", "Fase II constructiva: talleres participativos + panel Delphi (3 rondas); operacionalización y distribuciones a priori", "Modelo GDA validado · Artículo 2"),
    ("Año 3", "Fase III evaluativa: simulación Monte Carlo, índices de Sobol, pruebas de estrés; piloto acotado", "Resultados de evaluación · Artículo 3"),
    ("Año 4", "Integración, discusión, conclusión (Cap. 9); redacción y depósito de la tesis", "Tesis depositada y defensa"),
]:
    row = cron.add_row().cells
    for i, txt in enumerate([periodo, fase, prod]):
        row[i].text = txt
        row[i].paragraphs[0].runs[0].font.size = Pt(9)

out = r"C:\Users\LENOVO\Documents\tesis-gda\08_Capitulos\Tabla_de_Contenido_v5.docx"
doc.save(out)
print("OK:", out)
