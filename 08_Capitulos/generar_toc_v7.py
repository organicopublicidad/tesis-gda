"""
Genera Tabla_de_Contenido_v7.docx
BASE: v6 (correcciones de Daniel)
ADICIONES v7:
  1. §1.4.4 — Bioeconomía territorial como palanca transversal del VCO
  2. §4.5.4 — Fragilidad institucional en PDET: robustez como horizonte (antifragilidad contextualizada)
  3. §5.6 NUEVA — Bioeconomía territorial y diversificación bio-basada (insertar antes de actual 5.6)
       actual 5.6 → 5.7 | actual 5.7 → 5.8
  4. §6.4.7 — Diagnóstico de oportunidades bio-basadas
  5. Referencias Isabel Álvarez integradas en §3.5, §3.6, §5.4
  6. Anexo 6 — Compromisos I+D+i Conv. 975 Minciencias (Misión Bioeconomía y Territorio)
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3);  section.right_margin = Cm(2.5)

AZUL   = RGBColor(0x1F, 0x49, 0x7D)
VERDE  = RGBColor(0x1A, 0x6B, 0x3A)   # nuevas secciones PIIOM/bioeconomía
MORADO = RGBColor(0x5B, 0x2D, 0x82)   # antifragilidad / fragilidad
NARANJA= RGBColor(0xC0, 0x50, 0x00)   # Minciencias / productos
GRIS   = RGBColor(0x70, 0x70, 0x70)

def ph(doc, text, bold=False, size=10, color=None, before=3, after=1, indent=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)
    if indent: p.paragraph_format.left_indent = Cm(indent)
    if align:  p.alignment = align
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p

def cap(doc, num, title, nota=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"CAPÍTULO {num}: {title}")
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = AZUL
    if nota:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(nota); r2.italic = True; r2.font.size = Pt(9); r2.font.color.rgb = GRIS

def sec(doc, num, title, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{num}. {title}")
    r.bold = True; r.font.size = Pt(10)
    if color: r.font.color.rgb = color

def sub(doc, num, title, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{num}. {title}"); r.font.size = Pt(9)
    if color: r.font.color.rgb = color

def bloque(doc, items):
    """items = [(num, title, [subs], color_opt)]"""
    for item in items:
        num, title, subs = item[0], item[1], item[2]
        color = item[3] if len(item) > 3 else None
        sec(doc, num, title, color)
        for s in subs:
            scol = None
            if isinstance(s, tuple): s, scol = s
            sub(doc, s.split(".")[0]+"."+s.split(".")[1] if s[0].isdigit() else "",
                s.split(". ",1)[1] if ". " in s and s[0].isdigit() else s, scol)

def bloque2(doc, items):
    """items = [(num, title, [(sub_num, sub_title, color)], section_color)]"""
    for item in items:
        num, title, subs = item[0], item[1], item[2]
        scol = item[3] if len(item)>3 else None
        sec(doc, num, title, scol)
        for s in subs:
            if len(s)==3:
                sub(doc, s[0], s[1], s[2])
            else:
                sub(doc, s[0], s[1])

# ══════════════════════════════════════════════
# PORTADA
# ══════════════════════════════════════════════
ph(doc,"TABLA DE CONTENIDO — PLAN DE INVESTIGACIÓN",True,14,AZUL,0,8,align=WD_ALIGN_PARAGRAPH.CENTER)

for label,val in [
    ("Título:","Gobernanza Digital Asociativa: un modelo sociotécnico para la captura de valor en origen en organizaciones caficultoras de territorios PDET de Colombia"),
    ("Programa:","Doctorado TAPAS — ETSIAAB, Universidad Politécnica de Madrid"),
    ("Directora:","Dra. Sonia Benito Hernández"),
    ("Candidato:","Daniel Alejandro Pajoy Bastos"),
    ("Versión:","0.7 — 17 de junio de 2026"),
]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(label+" "); r.bold=True; r.font.size=Pt(11)
    r2=p.add_run(val); r2.font.size=Pt(11)

doc.add_paragraph()
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(10)
r=p.add_run("Pregunta de investigación: "); r.bold=True; r.font.size=Pt(11)
r2=p.add_run("¿Cómo un modelo sociotécnico de gobernanza digital asociativa habilita a las organizaciones caficultoras de pequeña escala en territorios PDET de Colombia para incrementar la proporción de valor capturado en origen (VCO) en la cadena global de valor del café?")
r2.italic=True; r2.font.size=Pt(11)
doc.add_paragraph("─"*80)

# ══════════════════════════════════════════════
# OBJETIVOS
# ══════════════════════════════════════════════
ph(doc,"OBJETIVOS DE LA INVESTIGACIÓN",True,12,AZUL,6,4)

p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
r=p.add_run("Objetivo general (OG): "); r.bold=True
p.add_run("Desarrollar y evaluar un modelo sociotécnico de gobernanza digital asociativa (GDA) que habilite a las organizaciones caficultoras de pequeña escala en territorios PDET de Colombia para incrementar la proporción de valor capturado en origen (VCO) en la cadena global de valor del café, en coherencia con la Misión Bioeconomía y Territorio (PIIOM — Conv. 975 Minciencias).")

oes=[
    ("OE1","I — Diagnóstica","Diagnosticar la distribución del valor a lo largo de la cadena del café colombiano, establecer la línea base VCO₀ y el capital social de las organizaciones en territorios PDET, construir el IPDET como moderador, identificar mediante fsQCA las configuraciones institucionales necesarias y suficientes para VCO alto, y mapear las oportunidades de diversificación bio-basada.","Cap. 6"),
    ("OE2","II — Constructiva","Diseñar de forma participativa, mediante Design Science Research y panel Delphi, el modelo GDA que articula el subsistema social (gobernanza asociativa, Ostrom) y el subsistema técnico (trazabilidad y soberanía de datos) bajo optimización conjunta, incorporando las palancas bioeconómicas como mecanismos transversales de VCO.","Cap. 7"),
    ("OE3","II → III","Operacionalizar y parametrizar los constructos del modelo (IG, componentes técnicos, VCO) y construir las distribuciones a priori, a partir del consenso Delphi, que alimentan la simulación.","Caps. 7-8"),
    ("OE4","III — Evaluativa","Evaluar el efecto del modelo GDA sobre el VCO y su robustez mediante Monte Carlo, índices de Sobol y escenarios adversos (pruebas de estrés), contrastando la hipótesis de que la gobernanza explica más varianza del VCO que la tecnología.","Cap. 8"),
]
t=doc.add_table(rows=1,cols=4); t.style="Table Grid"
for i,txt in enumerate(["OE","Fase","Enunciado","Capítulo"]):
    c=t.rows[0].cells[i]; c.text=txt
    c.paragraphs[0].runs[0].bold=True; c.paragraphs[0].runs[0].font.size=Pt(10)
for oe,fase,enun,cap_ in oes:
    row=t.add_row().cells
    for i,txt in enumerate([oe,fase,enun,cap_]):
        row[i].text=txt; row[i].paragraphs[0].runs[0].font.size=Pt(9)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8)
r=p.add_run("Hipótesis central: "); r.bold=True
p.add_run("La dimensión de gobernanza (subsistema social) explica una mayor proporción de la varianza del VCO que la dimensión tecnológica (subsistema técnico), evidenciando que la optimización conjunta (no solo lo digital) es la condición habilitante. Se contrasta con la descomposición de varianza de Sobol.")
p.paragraph_format.space_after=Pt(10)
doc.add_paragraph("─"*80)

# Prefacio
for txt in ["ABSTRACT (Inglés)","RESUMEN (Español)","LISTA DE TABLAS","LISTA DE FIGURAS","LISTA DE ACRÓNIMOS Y ABREVIATURAS"]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(txt); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x40,0x40,0x40)
doc.add_paragraph()

# ══════════════════════════════════════════════
# CAPÍTULO 1
# ══════════════════════════════════════════════
cap(doc,1,"INTRODUCCIÓN A LA TESIS")
bloque2(doc,[
    ("1.1","Introducción",[]),
    ("1.2","Contexto de la investigación",[
        ("1.2.1","El café en la economía colombiana",None),
        ("1.2.2","Distribución del valor en la cadena global y la brecha del VCO",None),
        ("1.2.3","El minifundio cafetero como universo de estudio",None),
    ]),
    ("1.3","El posconflicto colombiano y los PDET",[
        ("1.3.1","El Acuerdo de Paz de 2016",None),
        ("1.3.2","Los Programas de Desarrollo con Enfoque Territorial",None),
        ("1.3.3","La caficultura como vehículo de reincorporación",None),
    ]),
    ("1.4","La gobernanza como factor habilitante",[
        ("1.4.1","Gobernanza de comunes en organizaciones productoras",None),
        ("1.4.2","La transformación digital de las organizaciones rurales",None),
        ("1.4.3","La sostenibilidad como atributo de valor",None),
        ("1.4.4","La bioeconomía territorial y la diversificación bio-basada como palancas transversales del VCO",VERDE),   # NUEVO v7
    ],AZUL),
    ("1.5","Planteamiento del problema de investigación",[
        ("1.5.1","La brecha del valor capturado en origen (VCO)",None),
        ("1.5.2","Los cuatro déficits estructurales",None),
        ("1.5.3","Pregunta de investigación, hipótesis y objetivos",None),
    ]),
    ("1.6","Visión general de la metodología",[
        ("1.6.1","Design Science Research (Hevner et al., 2004): tres fases",None),
        ("1.6.2","Análisis Cualitativo Comparativo (fsQCA): censo FEMNCAFE (~20)",None),
        ("1.6.3","Simulación Monte Carlo e índices de sensibilidad de Sobol",None),
        ("1.6.4","Enfoque mixto secuencial: integración de métodos cualitativos y cuantitativos",None),
    ]),
    ("1.7","Contribuciones esperadas",[
        ("1.7.1","Contribución teórica",None),
        ("1.7.2","Contribución metodológica",None),
        ("1.7.3","Contribución praxeológica",None),
    ]),
    ("1.8","Estructura de la tesis",[]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 2
# ══════════════════════════════════════════════
cap(doc,2,"EL CAFÉ EN COLOMBIA Y EN EL MUNDO")
bloque2(doc,[
    ("2.1","Introducción",[]),
    ("2.2","Historia del café en Colombia",[]),
    ("2.3","Tendencias mundiales en la producción",[
        ("2.3.1","Principales países productores",None),
        ("2.3.2","Posición de Colombia en el mercado mundial",None),
        ("2.3.3","Evolución de la productividad",None),
    ]),
    ("2.4","Consumo, importaciones y exportaciones",[
        ("2.4.1","Principales mercados de destino del café colombiano",None),
        ("2.4.2","Estados Unidos como mercado principal",None),
        ("2.4.3","La Unión Europea como mercado estratégico",None),
    ]),
    ("2.5","Precios del café",[
        ("2.5.1","El Contrato C y la formación internacional de precios",None),
        ("2.5.2","El precio interno colombiano y el FEPCafé",None),
        ("2.5.3","Diferencial UGQ y prima de calidad",None),
    ]),
    ("2.6","La cadena de valor del café colombiano",[
        ("2.6.1","Eslabones productivos y comerciales",None),
        ("2.6.2","Distribución del valor entre actores",None),
        ("2.6.3","El caficultor en el eslabón de menor captura",None),
    ]),
    ("2.7","Segmentación del café",[
        ("2.7.1","Café convencional vs. cafés especiales",None),
        ("2.7.2","Métodos de procesamiento (lavado, honey, natural)",None),
        ("2.7.3","Café de origen y trazabilidad",None),
    ]),
    ("2.8","Políticas e institucionalidad cafetera en Colombia",[
        ("2.8.1","La Federación Nacional de Cafeteros (FNC)",None),
        ("2.8.2","El FEPCafé y el costo medio de producción",None),
        ("2.8.3","Programas de apoyo y financiamiento",None),
    ]),
    ("2.9","Caficultura y posconflicto",[
        ("2.9.1","Territorios PDET cafeteros",None),
        ("2.9.2","Iniciativas de café para la paz",None),
        ("2.9.3","La red FEMNCAFE",None),
    ]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 3
# ══════════════════════════════════════════════
cap(doc,3,"GOBERNANZA Y CADENAS GLOBALES DE VALOR")
bloque2(doc,[
    ("3.1","Introducción",[]),
    ("3.2","El marco analítico de las Cadenas Globales de Valor",[
        ("3.2.1","Origen y evolución del concepto",None),
        ("3.2.2","Las tres variables determinantes de Gereffi, Humphrey y Sturgeon",None),
        ("3.2.3","Los cinco tipos de gobernanza",None),
    ]),
    ("3.3","La cadena cautiva del café colombiano",[
        ("3.3.1","Caracterización como oligopsonio",None),
        ("3.3.2","Mecanismos de extracción de renta territorial",None),
        ("3.3.3","Asimetrías estructurales de poder",None),
    ]),
    ("3.4","La evolución del poder en la cadena del café (Grabs y Ponte, 2019)",[
        ("3.4.1","Tres fases históricas (pre-1989, 1989-2000, post-2000)",None),
        ("3.4.2","Los cuatro tipos de poder en la cadena",None),
        ("3.4.3","La paradoja de los sellos de sostenibilidad",None),
    ]),
    ("3.5","Upgrading en cadenas globales",[
        ("3.5.1","Upgrading de proceso, producto, funcional e intersectorial",None),
        ("3.5.2","Trayectorias posibles para organizaciones en países de ingreso medio (Álvarez & Marín, 2008)",None),   # Álvarez ref
        ("3.5.3","Hacia cadenas modulares y relacionales",None),
    ]),
    ("3.6","Redes de producción global (GPN 2.0)",[
        ("3.6.1","Actores no empresariales en la cadena",None),
        ("3.6.2","El papel del Estado, las ONG y la cooperación internacional",None),
        ("3.6.3","Aplicación al contexto PDET",None),
        ("3.6.4","Innovación en subsidiarias y cadenas colombianas: implicaciones para el café (Albis & Álvarez, 2008)",None),   # Álvarez ref
    ]),
    ("3.7","Diagnóstico del café colombiano como caso paradigmático",[]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 4
# ══════════════════════════════════════════════
cap(doc,4,"GOBERNANZA DE COMUNES, POSCONFLICTO Y SOSTENIBILIDAD")
bloque2(doc,[
    ("4.1","Introducción",[]),
    ("4.2","La teoría de gobernanza de comunes (Ostrom, 1990)",[
        ("4.2.1","El problema de los recursos de uso común",None),
        ("4.2.2","Los ocho principios de diseño institucional",None),
        ("4.2.3","Validación empírica multicontextual (Agrawal, 2001)",None),
    ]),
    ("4.3","Capital social y posconflicto en Colombia",[
        ("4.3.1","Erosión del capital social por el conflicto armado (Rettberg, 2010)",None),
        ("4.3.2","La cultura colectiva de los firmantes de paz como activo",None),
        ("4.3.3","Ausencia estatal histórica como causa estructural",None),
    ]),
    ("4.4","Cooperativismo agroalimentario y sus fallas (Ortmann y King, 2007)",[
        ("4.4.1","Estructuras de gobernanza en cooperativas cafeteras colombianas",None),
        ("4.4.2","Análisis institucional de la red FEMNCAFE",None),
    ]),
    ("4.5","Adaptación de Ostrom al posconflicto",[
        ("4.5.1","Reformulación contextual de los ocho principios",None),
        ("4.5.2","Principios complementarios para alta fragilidad institucional",None),
        ("4.5.3","Operacionalización: el Índice de Gobernanza (IG)",None),
        ("4.5.4","Fragilidad institucional en PDET: de la vulnerabilidad a la robustez como horizonte operativo",MORADO),   # NUEVO v7 — antifragilidad contextualizada
    ]),
    ("4.6","Investigación acción participativa (Fals Borda, 1986)",[
        ("4.6.1","Fundamento epistemológico para la co-construcción",None),
        ("4.6.2","Implicaciones metodológicas",None),
    ]),
    ("4.7","Sostenibilidad y compromiso ambiental en PYMES",[
        ("4.7.1","Apoyo público e inversión ambiental (Benito-Hernández et al., 2023)",None),
        ("4.7.2","PDET como instrumento de política habilitadora",None),
    ]),
    ("4.8","Valor implícito de los atributos de sostenibilidad en el café",[
        ("4.8.1","El modelo de precios hedónicos (Merbah & Benito-Hernández, 2023)",None),
        ("4.8.2","La prima por trazabilidad verificada",None),
        ("4.8.3","Implicaciones para el VCO",None),
    ]),
])

# Nota explicativa 4.5.4
p=doc.add_paragraph()
p.paragraph_format.left_indent=Cm(1.2); p.paragraph_format.space_after=Pt(4)
r=p.add_run("Nota §4.5.4: se define la antifragilidad (Taleb, 2012) como horizonte conceptual para caracterizar el grado de fragilidad institucional de las organizaciones PDET. La tesis no la opera como variable ni como eje metodológico — ese rol lo cumplen los escenarios adversos (pruebas de estrés, Cap. 8). La sección justifica por qué la robustez bajo incertidumbre es el proxy metodológico apropiado para este tipo de territorios.")
r.font.size=Pt(8); r.italic=True; r.font.color.rgb=MORADO

# ══════════════════════════════════════════════
# CAPÍTULO 5
# ══════════════════════════════════════════════
cap(doc,5,"TECNOLOGÍA DIGITAL, SOBERANÍA DE DATOS Y DISEÑO SOCIOTÉCNICO")
bloque2(doc,[
    ("5.1","Introducción",[]),
    ("5.2","La teoría de sistemas sociotécnicos",[
        ("5.2.1","Orígenes: la tradición de Tavistock (Trist y Bamforth, 1951)",None),
        ("5.2.2","Sistemas abiertos y optimización conjunta (Emery y Trist, 1965; Trist, 1981)",None),
        ("5.2.3","Vigencia contemporánea (Abbas y Michael, 2026; Horváth y Kemser, 2026)",None),
        ("5.2.4","El subsistema social y el subsistema técnico del GDA",None),
    ]),
    ("5.3","Blockchain y gobernanza de comunes",[
        ("5.3.1","Las seis affordances de Rozas et al. (2021)",None),
        ("5.3.2","Codificación de la confianza",None),
        ("5.3.3","Riesgos de la adopción top-down",None),
    ]),
    ("5.4","Blockchain en cadenas agroalimentarias",[
        ("5.4.1","Trazabilidad y verificación (Werbach, 2018)",None),
        ("5.4.2","Aplicaciones en café (Miatton y Amado, 2020; Elmatsani et al., 2026)",None),
        ("5.4.3","El EUDR y la georreferenciación obligatoria",None),
        ("5.4.4","Transformación digital en Iberoamérica y brecha digital rural (Álvarez, Quirós et al., 2021)",None),   # Álvarez ref
    ]),
    ("5.5","Soberanía de datos como recurso de uso común",[
        ("5.5.1","Los datos de la finca como activo colectivo",None),
        ("5.5.2","Arquitecturas técnicas de soberanía",None),
        ("5.5.3","Implicaciones éticas",None),
    ]),
    # ── NUEVA §5.6 Bioeconomía ─────────────────────────────────────────
    ("5.6","Bioeconomía territorial y diversificación bio-basada como palanca transversal del VCO",[
        ("5.6.1","Valorización de biomasa residual del café: pulpa, mucílago, cisco y posos (Solarte-Toro & Cardona, 2023)",VERDE),
        ("5.6.2","Agroecosistemas cafeteros en mercados voluntarios de carbono: protocolos MRV (González & Serna, 2018)",VERDE),
        ("5.6.3","Diversificación funcional: del grano a la biorrefinería de pequeña escala como estrategia de VCO",VERDE),
    ],VERDE),
    # ── §5.7 = antiguo §5.6 ───────────────────────────────────────────
    ("5.7","Instrumentos financieros para la captura de valor bio-basado: blended finance y Vehículos Especiales de Inversión (VEI)",[
        ("5.7.1","Blended finance para agricultura (Havemann et al., 2020)",None),
        ("5.7.2","VEI y capital ReFi: estructuras para diluir el riesgo rural (CAF, 2021)",None),
        ("5.7.3","Instrumentos por horizonte temporal e implicaciones para el VCO",None),
    ]),
    # ── §5.8 = antiguo §5.7 ───────────────────────────────────────────
    ("5.8","Síntesis del marco teórico integrado",[
        ("5.8.1","Articulación de los pilares: CGV + comunes + digital + bioeconomía bajo la lente sociotécnica",None),
        ("5.8.2","El modelo conceptual del GDA como sistema sociotécnico de optimización conjunta",None),
        ("5.8.3","Proposiciones teóricas e hipótesis derivadas",None),
    ]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 6
# ══════════════════════════════════════════════
cap(doc,6,"DIAGNÓSTICO DEL VALOR CAPTURADO EN ORIGEN — ANÁLISIS DE LA CADENA DEL CAFÉ COLOMBIANO PDET",
    "(Fase I — Diagnóstica y para la construcción del primer artículo)")
bloque2(doc,[
    ("6.1","Introducción",[]),
    ("6.2","Marco analítico de la distribución de valor",[
        ("6.2.1","Definición y medición del Valor Capturado en Origen (VCO)",None),
        ("6.2.2","Variables del modelo de diagnóstico",None),
        ("6.2.3","Construcción del Índice de Consolidación PDET (IPDET) como moderador",None),
    ]),
    ("6.3","Metodología",[
        ("6.3.1","Enfoque mixto secuencial y rol del DSR en la Fase I",None),
        ("6.3.2","Selección de organizaciones y territorios",None),
        ("6.3.3","Instrumentos cualitativos: entrevistas semiestructuradas, SNA y mapeo de la cadena",None),
        ("6.3.4","Análisis Cualitativo Comparativo (fsQCA): configuraciones institucionales para VCO alto",None),
        ("6.3.5","Integración de resultados cualitativos y cuantitativos",None),
    ]),
    ("6.4","Resultados",[
        ("6.4.1","Mapeo de la cadena del café colombiano hasta el mercado europeo",None),
        ("6.4.2","Distribución cuantitativa del valor por eslabón",None),
        ("6.4.3","VCO inicial (VCO₀) de las organizaciones estudiadas",None),
        ("6.4.4","Diagnóstico del capital social posconflicto (resultados SNA)",None),
        ("6.4.5","IPDET por municipio y categorización territorial",None),
        ("6.4.6","Configuraciones necesarias y suficientes para VCO alto (resultados fsQCA)",None),
        ("6.4.7","Diagnóstico de oportunidades bio-basadas: biomasa residual, servicios ecosistémicos y mercados voluntarios de carbono",VERDE),   # NUEVO v7
    ]),
    ("6.5","Discusión",[
        ("6.5.1","Comparación con la literatura previa",None),
        ("6.5.2","Patrones emergentes en la red FEMNCAFE",None),
        ("6.5.3","Implicaciones para el diseño del modelo GDA",None),
    ]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 7
# ══════════════════════════════════════════════
cap(doc,7,"DISEÑO Y VALIDACIÓN DEL MODELO GDA — ENFOQUE PARTICIPATIVO Y DELPHI",
    "(Fase II — Constructiva y para producción del segundo artículo)")
bloque2(doc,[
    ("7.1","Introducción",[]),
    ("7.2","Diseño científico de artefactos (Design Science Research)",[
        ("7.2.1","El paradigma de Hevner et al. (2004)",None),
        ("7.2.2","Las tres condiciones de validez",None),
        ("7.2.3","Aplicación al GDA",None),
    ]),
    ("7.3","Metodología",[
        ("7.3.1","Talleres de diseño participativo con las organizaciones",None),
        ("7.3.2","Panel Delphi: composición experta y protocolo de las tres rondas",None),
        ("7.3.3","Validación comunitaria",None),
    ]),
    ("7.4","Resultados",[
        ("7.4.1","Subsistema Social: Co-diseño de los mecanismos de gobernanza",None),
        ("7.4.2","Subsistema Técnico: Arquitectura técnica del componente blockchain",None),
        ("7.4.3","Consenso experto Delphi por variable",None),
        ("7.4.4","Distribuciones a priori construidas para la simulación",None),
        ("7.4.5","Prototipo del modelo GDA validado",None),
    ]),
    ("7.5","Discusión",[
        ("7.5.1","El modelo GDA como artefacto sociotécnico de optimización conjunta",None),
        ("7.5.2","Soberanía de datos como innovación institucional",None),
        ("7.5.3","Diferenciación frente a modelos previos",None),
    ]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 8
# ══════════════════════════════════════════════
cap(doc,8,"EVALUACIÓN DEL MODELO — MONTE CARLO, SOBOL Y PRUEBAS DE ESTRÉS",
    "(Fase III — Evaluativa para la elaboración del tercer artículo)")
bloque2(doc,[
    ("8.1","Introducción",[]),
    ("8.2","Marco evaluativo",[
        ("8.2.1","Contraste de hipótesis",None),
        ("8.2.2","Variables, parámetros e indicadores",None),
        ("8.2.3","Indicadores de robustez del VCO",None),
    ]),
    ("8.3","Metodología",[
        ("8.3.1","Configuración del modelo de simulación Monte Carlo",None),
        ("8.3.2","Implementación de los índices de sensibilidad de Sobol",None),
        ("8.3.3","Diseño de escenarios adversos",None),
        ("8.3.4","Piloto acotado e indicadores tempranos de adopción",None),
    ]),
    ("8.4","Resultados",[
        ("8.4.1","Distribución del VCO bajo el modelo GDA",None),
        ("8.4.2","Descomposición de varianza: gobernanza vs. tecnología",None),
        ("8.4.3","Respuesta del VCO bajo escenarios adversos",None),
        ("8.4.4","Robustez de los resultados",None),
        ("8.4.5","Resultados del piloto acotado",None),
        ("8.4.6","Transferibilidad en la red FEMNCAFE",None),
    ]),
    ("8.5","Discusión",[
        ("8.5.1","Validación o refutación de la hipótesis central",None),
        ("8.5.2","Alcance y límites de la robustez de los hallazgos",None),
        ("8.5.3","Limitaciones del modelo",None),
    ]),
])

# ══════════════════════════════════════════════
# CAPÍTULO 9
# ══════════════════════════════════════════════
cap(doc,9,"CONCLUSIÓN GENERAL")
bloque2(doc,[
    ("9.1","Síntesis de hallazgos por capítulo",[]),
    ("9.2","Respuesta a la pregunta de investigación",[]),
    ("9.3","Validación de las hipótesis y cumplimiento de los objetivos",[]),
    ("9.4","Contribuciones originales al conocimiento",[
        ("9.4.1","Aportes teóricos",None),
        ("9.4.2","Aportes metodológicos",None),
        ("9.4.3","Aportes praxeológicos",None),
    ]),
    ("9.5","Implicaciones para la política pública",[
        ("9.5.1","Para el marco PDET",None),
        ("9.5.2","Para la política cafetera colombiana",None),
        ("9.5.3","Para la regulación europea de sostenibilidad (EUDR)",None),
        ("9.5.4","Para la Misión Bioeconomía y Territorio (PIIOM — Conv. 975 Minciencias)",VERDE),
    ]),
    ("9.6","Implicaciones prácticas para las organizaciones",[]),
    ("9.7","Limitaciones del estudio",[]),
    ("9.8","Líneas futuras de investigación",[]),
])

# ══════════════════════════════════════════════
# ANEXOS
# ══════════════════════════════════════════════
doc.add_paragraph()
p=doc.add_paragraph(); r=p.add_run("ANEXOS")
r.bold=True; r.font.size=Pt(12); r.font.color.rgb=AZUL; p.paragraph_format.space_before=Pt(10)

anexos=[
    ("Anexo 1","Instrumentos de campo",[
        ("A1.1","Guion de entrevistas semiestructuradas",None),
        ("A1.2","Protocolo del Análisis de Redes Sociales (SNA)",None),
        ("A1.3","Cuestionario del panel Delphi",None),
        ("A1.4","Formularios de consentimiento informado",None),
        ("A1.5","Encuesta de transferibilidad FEMNCAFE",None),
    ]),
    ("Anexo 2","Operacionalización de constructos",[
        ("A2.1","Índice de Gobernanza (IG): indicadores y escalas de calibración",None),
        ("A2.2","IPDET: dimensiones e indicadores",None),
        ("A2.3","Cálculo del VCO: ecuaciones y fuentes de datos",None),
        ("A2.4","Tablas de calibración fsQCA",None),
    ]),
    ("Anexo 3","Implementación computacional",[
        ("A3.1","Código de la simulación Monte Carlo",None),
        ("A3.2","Implementación de los índices de Sobol",None),
        ("A3.3","Diseño de escenarios adversos",None),
        ("A3.4","Validación del modelo",None),
    ]),
    ("Anexo 4","Aprobación ética",[
        ("A4.1","Dictamen del comité de ética institucional",None),
        ("A4.2","Convenios con organizaciones participantes",None),
    ]),
    ("Anexo 5","Publicaciones derivadas",[
        ("A5.1","Artículo 1: Diagnóstico de la distribución de valor (Fase I)",None),
        ("A5.2","Artículo 2: El modelo de gobernanza GDA (Fase II)",None),
        ("A5.3","Artículo 3: Evaluación bajo incertidumbre (Fase III)",None),
        ("A5.4","Ponencia en congreso internacional de economía agraria, gobernanza digital o bioeconomía",None),
    ]),
    # ── NUEVO Anexo 6 — Minciencias ───────────────────────────────────
    ("Anexo 6","Compromisos I+D+i — Convocatoria 975 Minciencias (Misión Bioeconomía y Territorio)",[
        ("A6.1","IAT-Café: especificación del Instrumento de Diagnóstico de Fragilidad/Robustez Territorial (producto Tipo B — derivado paralelo a la tesis)",NARANJA),
        ("A6.2","Arquitectura GDA: documento de diseño de acceso abierto (producto Tipo B — especificación técnica transferible)",NARANJA),
        ("A6.3","Protocolo MRV para mercados voluntarios de carbono en agroecosistemas cafeteros (producto Tipo B)",NARANJA),
        ("A6.4","Manual de transferencia técnica y piloto de formación en organización del Cauca (producto Tipo C)",NARANJA),
        ("A6.5","Documento de política pública — lineamientos para Minciencias, MinAgricultura, FNC y ART (producto Tipo C — armonizado con §9.5.4)",NARANJA),
        ("A6.6","Protocolo de taller de co-creación: metodología de facilitación documentada (producto Tipo C — acceso abierto)",NARANJA),
    ],NARANJA),
]
for anx in anexos:
    label,title,subs = anx[0],anx[1],anx[2]
    hcol = anx[3] if len(anx)>3 else None
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(f"{label}. {title}"); r.bold=True; r.font.size=Pt(10)
    if hcol: r.font.color.rgb=hcol
    for s in subs:
        p2=doc.add_paragraph(); p2.paragraph_format.left_indent=Cm(0.7)
        p2.paragraph_format.space_before=Pt(0); p2.paragraph_format.space_after=Pt(0)
        r2=p2.add_run(f"{s[0]}. {s[1]}"); r2.font.size=Pt(9)
        if s[2]: r2.font.color.rgb=s[2]

# ══════════════════════════════════════════════
# REFERENCIAS + RD 99/2011
# ══════════════════════════════════════════════
doc.add_paragraph()
p=doc.add_paragraph(); r=p.add_run("REFERENCIAS"); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=AZUL
doc.add_paragraph("(En orden alfabético, formato APA 7, con DOI cuando esté disponible)").runs[0].italic=True
doc.add_paragraph("─"*80)

p=doc.add_paragraph(); r=p.add_run("MEDIOS Y RECURSOS (RD 99/2011)")
r.bold=True; r.font.size=Pt(11); r.font.color.rgb=AZUL; p.paragraph_format.space_before=Pt(10)
medios=[
    "Acceso al terreno: vínculo operativo con COFFIS y la red FEMNCAFE (~20 organizaciones — censo completo). Sede del candidato en Popayán, Cauca (caso de validación).",
    "Datos: entrevistas semiestructuradas, datos de la cadena por eslabón, datos de finca para trazabilidad; panel de expertos para el Delphi; condiciones institucionales para el fsQCA; datos de biomasa residual y servicios ecosistémicos para el diagnóstico bio-basado.",
    "Recursos computacionales: Python (Monte Carlo, Sobol, 10.000 iteraciones); software fsQCA 3.0 / R-QCA; herramientas SNA.",
    "Dirección y marco institucional: Dra. Sonia Benito (ETSIAAB-UPM, grupo BIDA); programa TAPAS; modalidad virtual desde Colombia con trabajo de campo in situ.",
    "Ética: dictamen del comité de ética institucional y convenios de consentimiento informado (Anexo 4).",
]
for m in medios:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    r=p.add_run("• "+m); r.font.size=Pt(10)

p=doc.add_paragraph(); r=p.add_run("PLANIFICACIÓN TEMPORAL (3–4 años)")
r.bold=True; r.font.size=Pt(11); r.font.color.rgb=AZUL; p.paragraph_format.space_before=Pt(8)
ct=doc.add_table(rows=1,cols=3); ct.style="Table Grid"
for i,txt in enumerate(["Periodo","Fase / Hito","Producto"]):
    c=ct.rows[0].cells[i]; c.text=txt
    c.paragraphs[0].runs[0].bold=True; c.paragraphs[0].runs[0].font.size=Pt(10)
for per,fase,prod in [
    ("Año 1","Marco teórico (Caps. 1-5) + Fase I diagnóstica: trabajo de campo, SNA, fsQCA, diagnóstico bio-basado, VCO₀, IPDET","CAPD · Artículo 1"),
    ("Año 2","Fase II constructiva: talleres participativos + panel Delphi (3 rondas); operacionalización y distribuciones a priori; especificación arquitectura GDA y protocolo MRV","Modelo GDA validado · Artículo 2 · Productos Minciencias Tipo B"),
    ("Año 3","Fase III evaluativa: Monte Carlo, Sobol, pruebas de estrés; piloto acotado; manual de transferencia; documento de política pública","Evaluación · Artículo 3 · Productos Minciencias Tipo C"),
    ("Año 4","Integración, conclusión (Cap. 9); redacción y depósito; ponencia internacional","Tesis depositada · Ponencia Tipo A"),
]:
    row=ct.add_row().cells
    for i,txt in enumerate([per,fase,prod]):
        row[i].text=txt; row[i].paragraphs[0].runs[0].font.size=Pt(9)

# ══════════════════════════════════════════════
# LEYENDA DE COLORES
# ══════════════════════════════════════════════
doc.add_paragraph()
p=doc.add_paragraph(); r=p.add_run("LEYENDA DE SECCIONES NUEVAS (v7)")
r.bold=True; r.font.size=Pt(9); r.font.color.rgb=GRIS
leyenda=[
    (VERDE,"Verde — Bioeconomía / diversificación bio-basada / PIIOM (secciones nuevas v7)"),
    (MORADO,"Morado — Fragilidad institucional / antifragilidad como contexto teórico (§4.5.4)"),
    (NARANJA,"Naranja — Compromisos I+D+i Minciencias Conv. 975 (Anexo 6)"),
]
for col,txt in leyenda:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1)
    r=p.add_run("  ● "+txt); r.font.size=Pt(8); r.font.color.rgb=col

out=r"C:\Users\LENOVO\Documents\tesis-gda\08_Capitulos\Tabla_de_Contenido_v7.docx"
doc.save(out)
print("OK:",out)
